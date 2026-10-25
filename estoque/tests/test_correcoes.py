"""Casos da revisão do conversor: cada teste reproduz uma entrada que antes gerava valor errado,
perdia item sem aviso ou saía num arquivo que o PVA recusa. Todo SPED gerado aqui passa pela
conferência do leiaute (validar), que segue o Guia Prático EFD ICMS/IPI."""

import os
import re
import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal
from pathlib import Path
from unittest import mock

APP = Path(__file__).resolve().parents[1]
ARQ = Path(__file__).resolve().parent / "arquivos"
TMP = Path(tempfile.mkdtemp(prefix="estoque-correcoes-"))
os.environ.setdefault("ESTOQUE_DADOS", str(TMP / "dados"))
sys.path.insert(0, str(APP))

import gerar_sped as gs  # noqa: E402

# Quantidade de campos de cada registro (REG incluído), pelo Guia Prático
LEIAUTE = {
    "0000": 15, "0001": 2, "0002": 2, "0005": 10, "0100": 14, "0150": 13, "0190": 3, "0200": 13, "0990": 2,
    "B001": 2, "B990": 2, "C001": 2, "C990": 2, "D001": 2, "D990": 2, "E001": 2, "E100": 3, "E110": 15, "E990": 2,
    "G001": 2, "G990": 2, "H001": 2, "H005": 4, "H010": 11, "H990": 2, "K001": 2, "K990": 2,
    "1001": 2, "1010": 14, "1990": 2, "9001": 2, "9900": 3, "9990": 2, "9999": 2,
}
N2, N3, N6 = re.compile(r"\d+,\d{2}"), re.compile(r"\d+,\d{3}"), re.compile(r"\d+,\d{6}")


def validar(teste, linhas):
    """Confere o arquivo como o PVA: leiaute de cada registro, totalizadores, formatos e as
    amarrações entre registros (0190 x unidades, 0200 x H010, 0150 x COD_PART)."""
    reg = [l[1:-1].split("|") for l in linhas]
    for l, c in zip(linhas, reg):
        teste.assertTrue(l.startswith("|") and l.endswith("|"), l)
        teste.assertIn(c[0], LEIAUTE, l)
        teste.assertEqual(len(c), LEIAUTE[c[0]], f"quantidade de campos: {l}")
        l.encode("latin-1")  # ISO-8859-1
    por = lambda r: [c for c in reg if c[0] == r]
    z = por("0000")[0]
    ini, fim = gs.data_ddmmaaaa(z[3]), gs.data_ddmmaaaa(z[4])
    teste.assertEqual(ini.day, 1)
    teste.assertEqual((fim.month, fim.day), (ini.month, gs.calendar.monthrange(fim.year, fim.month)[1]))
    teste.assertEqual(z[1], gs.VERSAO_LEIAUTE.get(fim.year, z[1]))
    # totalizadores
    teste.assertEqual(int(por("0990")[0][1]), sum(c[0][0] == "0" for c in reg))
    teste.assertEqual(int(por("H990")[0][1]), sum(c[0][0] == "H" for c in reg))
    teste.assertEqual(int(por("9990")[0][1]), sum(c[0][0] == "9" for c in reg))
    teste.assertEqual(int(por("9999")[0][1]), len(reg))
    cont = {c[1]: int(c[2]) for c in por("9900")}
    teste.assertEqual(set(cont), {c[0] for c in reg})
    for r, n in cont.items():
        teste.assertEqual(n, len(por(r)), r)
    # inventário
    h005, h010 = por("H005"), por("H010")
    teste.assertEqual(len(h005), 1)
    teste.assertLessEqual(gs.data_ddmmaaaa(h005[0][1]), fim)
    teste.assertEqual(h005[0][3], "01")
    teste.assertRegex(h005[0][2], N2)
    teste.assertEqual(gs.num(h005[0][2]), sum(gs.num(c[5]) for c in h010))
    unidades, cad = {c[1] for c in por("0190")}, {c[1]: c for c in por("0200")}
    parts = {c[1] for c in por("0150")}
    teste.assertEqual(len(cad), len(por("0200")), "0200 repetido")
    for c in h010:
        teste.assertIn(c[1], cad)
        teste.assertLessEqual(len(c[1]), 60)
        teste.assertIn(c[2], unidades)
        teste.assertEqual(c[2], cad[c[1]][5], "unidade do H010 = unidade do 0200")
        teste.assertRegex(c[3], N3)
        teste.assertRegex(c[4], N6)
        teste.assertRegex(c[5], N2)
        teste.assertRegex(c[10], N2)
        teste.assertGreater(gs.num(c[3]), 0)
        teste.assertLessEqual(gs.num(c[10]), gs.num(c[5]))
        teste.assertIn(c[6], ("0", "1", "2"))
        if c[6] == "0":
            teste.assertEqual(c[7], "")
        else:
            teste.assertIn(c[7], parts, "COD_PART do H010 sem 0150")
    for c in cad.values():
        teste.assertIn(c[6], gs.TIPOS_ITEM)
        teste.assertIn(len(c[7]), (0, 8))
        teste.assertIn(len(c[12]), (0, 7))
        teste.assertIn(c[5], unidades)
    for c in por("0150"):
        teste.assertTrue(c[2] and c[3] and (len(c[7]) == 7 or c[3].lstrip("0") != "1058"))


def arquivo(nome, conteudo, encoding="utf-8"):
    p = TMP / nome
    if isinstance(conteudo, bytes):
        p.write_bytes(conteudo)
    else:
        p.write_text(conteudo, encoding=encoding)
    return p


class Base(unittest.TestCase):
    def gerar(self, *caminhos, **opcoes):
        fontes = [gs.ler_arquivo(Path(c)) for c in caminhos]
        principal = gs.escolher_principal(fontes)
        opcoes.setdefault("data_inv", "31/12/2025")
        res = gs.converter(principal, [f for f in fontes if f is not principal], gs.opcoes(sem_consulta=True, **opcoes))
        validar(self, res.registros)
        return res

    def reg(self, res, r):
        return [l[1:-1].split("|") for l in res.registros if l.startswith(f"|{r}|")]

    def h010(self, res):
        return {c[1]: c for c in self.reg(res, "H010")}

    def pend(self, res, trecho, tipo=None):
        return [p for p in res.pendencias if trecho in p[2] and (tipo is None or p[0] == tipo)]


class Numeros(unittest.TestCase):
    """Itens 1, 2, 25 e 26: o separador decimal e os formatos que o Excel e os ERPs gravam."""

    def test_virgula_e_ponto_o_ultimo_e_o_decimal(self):
        self.assertEqual(gs.num("1,234.56"), Decimal("1234.56"))
        self.assertEqual(gs.num("1.234,56"), Decimal("1234.56"))
        self.assertEqual(gs.num("1,800,600.00"), Decimal("1800600.00"))
        self.assertEqual(gs.num("1.800.600,00"), Decimal("1800600.00"))
        self.assertEqual(gs.num("R$ 1.234,56"), Decimal("1234.56"))

    def test_um_separador_so(self):
        self.assertEqual(gs.num("2,5"), Decimal("2.5"))
        self.assertEqual(gs.num("2.5"), Decimal("2.5"))
        self.assertEqual(gs.num("1.500"), Decimal("1.500"))  # sem pista do arquivo: decimal
        self.assertEqual(gs.num("1.500", br=True), Decimal("1500"))
        self.assertEqual(gs.num("1,500", en=True), Decimal("1500"))
        self.assertEqual(gs.num("1,5", en=True), Decimal("1.5"))  # não é milhar
        self.assertEqual(gs.num("1.234.567"), Decimal("1234567"))
        self.assertEqual(gs.num("1,234,567"), Decimal("1234567"))

    def test_negativo_contabil(self):
        self.assertEqual(gs.num("(1.234,56)"), Decimal("-1234.56"))
        self.assertEqual(gs.num("1.234,56-"), Decimal("-1234.56"))
        self.assertIsNone(gs.num("()"))
        self.assertIsNone(gs.num("-"))

    def test_celulas_que_nao_sao_numero(self):
        self.assertIsNone(gs.num(True))  # antes: InvalidOperation derrubava a planilha
        self.assertIsNone(gs.num(False))
        self.assertIsNone(gs.num(float("nan")))
        self.assertIsNone(gs.num(float("inf")))
        self.assertIsNone(gs.num("NaN"))
        self.assertIsNone(gs.num("Infinity"))
        self.assertIsNone(gs.num("abc"))
        self.assertEqual(gs.num(10), Decimal(10))

    def test_estilo_do_arquivo(self):
        self.assertEqual(gs.estilo_numeros(["1.500", "2,50"], "en"), ("br", ""))
        self.assertEqual(gs.estilo_numeros(["1,500", "2.50"], "br"), ("en", ""))
        self.assertEqual(gs.estilo_numeros(["0,500"], "en")[0], "br")  # 0,500 não é milhar
        estilo, aviso = gs.estilo_numeros(["1.500", "3.000"], "br")
        self.assertEqual(estilo, "br")
        self.assertIn("1500", aviso)
        self.assertIn("mistura", gs.estilo_numeros(["1.234,56", "1,234.56"], "br")[1])


class Planilhas(Base):
    def test_csv_ingles_com_milhar(self):
        # Antes: 1.800.600,00 virava 1,80, sem aviso
        p = arquivo("en.csv", 'SKU,Description,Unit,Qty,Unit Cost,Total\nK1,CABO,UN,"1,200","1,500.50","1,800,600.00"\n')
        h = self.h010(self.gerar(p))
        self.assertEqual(h["K1"][3:6], ["1200,000", "1500,500000", "1800600,00"])

    def test_csv_brasileiro_so_com_inteiros(self):
        # Antes: "1.500" virava 1,5 e "3.000" virava 3, sem aviso
        p = arquivo("br_int.csv", "Codigo;Descricao;Qtd;Valor Total\nA1;CANETA;1.500;3.000\n")
        res = self.gerar(p)
        self.assertEqual(self.h010(res)["A1"][3:6], ["1500,000", "2,000000", "3000,00"])
        self.assertTrue(self.pend(res, "podem ser milhar ou decimal", "arquivo"))

    def test_csv_com_virgula_so_com_inteiros(self):
        p = arquivo("en_int.csv", 'Code,Description,Qty,Total\nA1,PEN,"1,500","3,000"\n')
        res = self.gerar(p)
        self.assertEqual(self.h010(res)["A1"][3], "1500,000")
        self.assertTrue(self.pend(res, "formato inglês"))

    def test_csv_ponto_e_virgula_com_muitas_virgulas_decimais(self):
        # Mais vírgulas (decimais) que pontos e vírgulas: o separador vem do cabeçalho
        p = arquivo("virgulas.csv", "Codigo;Descricao;Qtd;Valor Unit\nA1;X, Y, Z;1,5;2,25\nA2;W, V;2,5;1,75\n")
        h = self.h010(self.gerar(p))
        self.assertEqual((h["A1"][3], h["A2"][4]), ("1,500", "1,750000"))

    def test_linha_sem_codigo_entra_com_codigo_gerado(self):
        # Antes: a linha sumia do SPED sem aviso
        p = arquivo("semcod.csv", "Codigo;Descricao;Qtd;Valor Unit\nA1;CANETA;10;2,00\n;LAPIS SEM CODIGO;5;1,00\n")
        res = self.gerar(p)
        h = self.h010(res)
        self.assertEqual(h["SEMCOD0001"][3:6], ["5,000", "1,000000", "5,00"])
        self.assertTrue(self.pend(res, "linha sem código"))
        self.assertEqual(res.total, Decimal("25.00"))

    def test_linha_sem_codigo_e_sem_descricao_com_numeros_e_avisada(self):
        p = arquivo("subtotal.csv", "Codigo;Descricao;Qtd;Valor Unit;Valor Total\nA1;CANETA;10;2,00;20,00\n;;10;;20,00\n")
        res = self.gerar(p)
        self.assertEqual(list(self.h010(res)), ["A1"])
        self.assertTrue(self.pend(res, "linha 3 sem código nem descrição"))

    def test_titulo_de_grupo_sem_numeros_nao_e_item(self):
        p = arquivo("grupo.csv", "Codigo;Descricao;Qtd;Valor Unit\n;BEBIDAS;;\nA1;REFRI;1;5,00\n")
        res = self.gerar(p)
        self.assertEqual(list(self.h010(res)), ["A1"])
        self.assertFalse(self.pend(res, "SEMCOD"))

    def test_xls_com_ncm_cest_e_tipo_numericos(self):
        # O .xls guarda número como decimal: 94036000.0 virava "940360000" (NCM apagado) e
        # 1022990.0 virava "10229900" (NCM errado, sem aviso); o tipo 1.0 virava "10"
        res = self.gerar(ARQ / "ncm_numerico.xls")
        cad = {c[1]: c for c in self.reg(res, "0200")}
        self.assertEqual((cad["101"][7], cad["101"][12], cad["101"][6]), ("94036000", "2803600", "00"))
        self.assertEqual((cad["102"][7], cad["102"][6]), ("01022990", "01"))
        self.assertEqual(cad["103"][7], "10063021")
        self.assertFalse(self.pend(res, "NCM"))

    def test_quantidade_em_terceiros_sem_participante_fica_fora_com_aviso(self):
        # Antes: os 7 em terceiros sumiam sem aviso
        p = arquivo("terc.csv", "Codigo;Descricao;Qtd. no estab;Qtd. em terceiros;Valor Unit\nA1;CANETA;10;7;2,00\n")
        res = self.gerar(p)
        self.assertEqual([(c[1], c[3], c[6]) for c in self.reg(res, "H010")], [("A1", "10,000", "0")])
        self.assertTrue(self.pend(res, "sem o participante", "item"))
        self.assertEqual([(i.codigo, i.qtd, i.ind_prop) for i in res.fora], [("A1", Decimal(7), "1")])

    def test_quantidade_em_terceiros_com_cnpj_gera_h010_e_0150(self):
        consulta = {"nome": "DEPOSITO GERAL LTDA", "uf": "CE", "cod_mun": "2304400", "cep": "60000000",
                    "end": "RUA B", "num": "20", "compl": "", "bairro": "CENTRO", "fantasia": "", "fone": "", "email": ""}
        p = arquivo("terc_cnpj.csv", "Codigo;Descricao;Qtd. no estab;Qtd. em terceiros;Valor Unit;CNPJ do terceiro\n"
                                     "A1;CANETA;10;7;2,00;11.444.777/0001-61\nA2;LAPIS;0;3;1,00;11.444.777/0001-61\n")
        fontes = [gs.ler_arquivo(p)]
        with mock.patch.object(gs, "consultar_cnpj", return_value=consulta) as c:
            res = gs.converter(fontes[0], [], gs.opcoes(data_inv="31/12/2025", sem_consulta=False, cnpj="11222333000181",
                                                         uf="CE", cod_mun="2304400", cep="60000000", end="X", bairro="Y"))
        self.assertIn(mock.call("11444777000161"), c.call_args_list)
        validar(self, res.registros)
        h = [(c[1], c[3], c[6], c[7]) for c in self.reg(res, "H010")]
        self.assertEqual(h, [("A1", "10,000", "0", ""), ("A1", "7,000", "1", "11444777000161"),
                             ("A2", "3,000", "1", "11444777000161")])
        self.assertEqual(self.reg(res, "0150"), [["0150", "11444777000161", "DEPOSITO GERAL LTDA", "1058", "11444777000161",
                                                  "", "", "2304400", "", "RUA B", "20", "", "CENTRO"]])
        self.assertEqual(len(self.reg(res, "0200")), 2)
        self.assertEqual(res.total, Decimal("37.00"))

    def test_aliquota_de_icms_nao_e_icms_recuperavel(self):
        # Antes: "Alíquota ICMS" 18 virava ICMS recuperável e o VL_ITEM_IR saía 2,00
        p = arquivo("aliq.csv", "Codigo;Descricao;Qtd;Valor Unit;Aliquota ICMS;Base ICMS\nA1;CANETA;10;2,00;18;20,00\n")
        self.assertEqual(self.h010(self.gerar(p))["A1"][10], "20,00")

    def test_icms_recuperavel_maior_que_o_item_e_ignorado(self):
        p = arquivo("icms.csv", "Codigo;Descricao;Qtd;Valor Unit;ICMS Recuperavel\nA1;CANETA;10;2,00;50,00\nA2;LAPIS;1;10,00;1,80\n")
        res = self.gerar(p)
        h = self.h010(res)
        self.assertEqual((h["A1"][10], h["A2"][10]), ("20,00", "8,20"))
        self.assertTrue(self.pend(res, "maior que o valor do item"))

    def test_outra_aba_com_inventario_e_avisada(self):
        import openpyxl
        caminho = TMP / "duas_abas.xlsx"
        wb = openpyxl.Workbook()
        for nome, cod in (("Loja 1", "L1"), ("Loja 2", "L2")):
            ws = wb.create_sheet(nome)
            ws.append(["Código", "Descrição", "Quantidade", "Valor Unitário"])
            ws.append([cod, "ITEM", 1, 10])
        wb.remove(wb["Sheet"])
        wb.save(caminho)
        res = self.gerar(caminho)
        self.assertEqual(list(self.h010(res)), ["L1"])
        self.assertTrue(self.pend(res, "a aba \"Loja 2\" também tem um inventário"))

    def test_aba_de_cadastro_antes_do_inventario(self):
        import openpyxl
        caminho = TMP / "cadastro_antes.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Cadastro"
        ws.append(["Código", "Descrição", "NCM"])
        ws.append(["C1", "CANETA", "96081000"])
        ws = wb.create_sheet("Inventário")
        ws.append(["Código", "Descrição", "Quantidade", "Valor Unitário"])
        ws.append(["C1", "CANETA", 4, 2.5])
        wb.save(caminho)
        self.assertEqual(self.h010(self.gerar(caminho))["C1"][5], "10,00")

    def test_csv_com_quebra_de_linha_dentro_de_aspas(self):
        p = arquivo("aspas.csv", 'Codigo;Descricao;Qtd;Valor Unit\nA1;"CANETA\nAZUL";1;2,00\nA2;LAPIS;1;1,00\n')
        res = self.gerar(p)
        self.assertEqual(self.reg(res, "0200")[0][2], "CANETA AZUL")
        self.assertEqual(len(self.reg(res, "H010")), 2)

    def test_celula_booleana_nao_derruba_a_planilha(self):
        import openpyxl
        caminho = TMP / "bool.xlsx"
        wb = openpyxl.Workbook()
        wb.active.append(["Código", "Descrição", "Quantidade", "Valor Unitário"])
        wb.active.append(["B1", "ITEM", True, 2])
        wb.active.append(["B2", "ITEM 2", 1, 2])
        wb.save(caminho)
        res = self.gerar(caminho)
        self.assertEqual(list(self.h010(res)), ["B2"])
        self.assertTrue(self.pend(res, "sem quantidade"))

    def test_negativo_contabil_fica_fora_com_o_motivo_certo(self):
        p = arquivo("contabil.csv", "Codigo;Descricao;Qtd;Valor Unit\nA1;X;2;(10,00)\nA2;Y;1;5,00\n")
        res = self.gerar(p)
        self.assertEqual(list(self.h010(res)), ["A2"])
        self.assertTrue(any(c == "A1" and "negativo" in m for _, c, m in res.pendencias))

    def test_indicador_de_propriedade_na_planilha(self):
        p = arquivo("prop.csv", "Codigo;Descricao;Qtd;Valor Unit;Ind. Prop\nA1;X;2;1,00;0\nA2;Y;1;5,00;X\n")
        res = self.gerar(p)
        self.assertEqual({c[1]: c[6] for c in self.reg(res, "H010")}, {"A1": "0", "A2": "0"})
        self.assertTrue(self.pend(res, "não é 0, 1 ou 2"))


class Relatorio(Base):
    CAB = "REGISTRO DE INVENTARIO\nFirma: LOJA\nEstoque existente em 31/12/2025\n"

    def test_linha_sem_ncm_entra(self):
        # Antes: a linha sem NCM sumia sem aviso
        rel = self.CAB + ("94036000 U3424 CADEIRA PRETA     UN      10,000      150,0000      1.500,00\n"
                          "         X55 ITEM SEM NCM        UN       2,000       10,0000         20,00\n")
        res = self.gerar(arquivo("rel_sem_ncm.txt", rel))
        h = self.h010(res)
        self.assertEqual(h["X55"][3:6], ["2,000", "10,000000", "20,00"])
        self.assertEqual(res.total, Decimal("1520.00"))

    def test_linha_com_tres_numeros_que_nao_casa_vira_pendencia(self):
        rel = self.CAB + ("94036000 U3424 CADEIRA PRETA     UN      10,000      150,0000      1.500,00\n"
                          "         ITEM COM UNIDADE COMPRIDA UNIDADES  2,000  10,0000  20,00\n")
        res = self.gerar(arquivo("rel_quebrado2.txt", rel))
        self.assertEqual(len(self.reg(res, "H010")), 1)
        self.assertTrue(self.pend(res, "linha não lida", "arquivo"))

    def test_rotulo_de_total_com_tres_numeros_nao_e_item(self):
        rel = self.CAB + ("94036000 U3424 CADEIRA PRETA     UN      10,000      150,0000      1.500,00\n"
                          "         TOTAL DO GRUPO          UN      10,000      150,0000      1.500,00\n")
        res = self.gerar(arquivo("rel_total.txt", rel))
        self.assertEqual(list(self.h010(res)), ["U3424"])
        # mas um produto com "TOTAL" no nome continua sendo item
        rel = self.CAB + "         TOTAL FLEX 1L           UN       1,000       30,0000         30,00\n"
        self.assertEqual(self.reg(self.gerar(arquivo("rel_total2.txt", rel)), "0200")[0][2], "TOTAL FLEX 1L")

    def test_codigo_com_espaco_vira_um_aviso_so(self):
        # Num relatório em que quase todo código tem espaço ("U0464 013/20MM"), um aviso por item
        # escondia os que precisam de conferência: agora vai um aviso do arquivo
        rel = self.CAB + ("91139000 U0464 013/20MM ADONE METAL  UN   1,000   10,0000   10,00\n"
                          "91139000 X100 50 CADEIRAS           UN   1,000   10,0000   10,00\n")
        res = self.gerar(arquivo("rel_cod.txt", rel))
        self.assertEqual(sorted(self.h010(res)), ["U0464 013/20MM", "X100 50"])
        self.assertFalse(self.pend(res, "separado da descrição", "item"))
        avisos = self.pend(res, "2 código(s) com espaço", "arquivo")
        self.assertEqual(len(avisos), 1)
        self.assertIn("U0464 013/20MM", avisos[0][2])

    def test_separacao_do_codigo_em_linhas_reais(self):
        casos = {
            # "P/" de "para" é descrição, mesmo com barra
            "PU11599JS/02 PULS.P/REL. PHALANX": ("PU11599JS/02", "PULS.P/REL. PHALANX", False),
            "CO13595JSB/13 2H COROA P/ REL. POLICE": ("CO13595JSB/13 2H", "COROA P/ REL. POLICE", False),
            # código grudado numa palavra no próprio relatório: separa, mas avisa no item
            "CX2038GSG/01CAIXA P/ REL. GARRIDO &": ("CX2038GSG/01CAIXA", "P/ REL. GARRIDO &", True),
            "WM1 0301BRINCOS MISS SIXTY": ("WM1 0301BRINCOS", "MISS SIXTY", True),
            # sem número: o código vem antes do " - " (suposição, avisada no item)
            "BIP LITE BLACK - SMARTWATCH AMAZFIT": ("BIP LITE BLACK", "SMARTWATCH AMAZFIT", True),
            "GG 699 - EXPOSITOR INDIVIDUAL": ("GG 699", "EXPOSITOR INDIVIDUAL", True),
            # termo curto com hífen
            "MW-CX CAIXA EM ACO P/ RELOGIO": ("MW-CX", "CAIXA EM ACO P/ RELOGIO", False),
            # os casos que já funcionavam continuam iguais
            "5095 - EXPOSITOR 6 BALDAS": ("5095", "EXPOSITOR 6 BALDAS", False),
            "5102- ISOTONIC DRINK": ("5102", "ISOTONIC DRINK", False),
            "BG 99100-003-A EXP MEDIO EM PLASTICO P/": ("BG 99100-003-A", "EXP MEDIO EM PLASTICO P/", False),
            "GC 1999885 EXPOSITOR EM MADEIRA": ("GC 1999885", "EXPOSITOR EM MADEIRA", False),
            "PA12173JSBNR/65 PARAFUSO": ("PA12173JSBNR/65", "PARAFUSO", False),
            # sem pista segura: continua sem código (vira SEMCOD com pendência)
            "KIT 2 PEÇAS": ("", "KIT 2 PEÇAS", False),
            "IPM 80 ESTOJO EM MADEIRA": ("", "IPM 80 ESTOJO EM MADEIRA", False),
            "SQUEEZE BRANCA 790ML": ("", "SQUEEZE BRANCA 790ML", False),
            "UMA DESCRICAO MUITO COMPRIDA DE PRODUTO - COM TRACO": ("", "UMA DESCRICAO MUITO COMPRIDA DE PRODUTO - COM TRACO", False),
        }
        for linha, esperado in casos.items():
            with self.subTest(linha):
                self.assertEqual(gs.separar_codigo(linha), esperado)

    def test_codigo_antes_do_traco_vai_para_as_pendencias(self):
        rel = self.CAB + "85176277 BIP LITE BLACK - SMARTWATCH  UN   1,000   10,0000   10,00\n"
        res = self.gerar(arquivo("rel_traco.txt", rel))
        self.assertEqual(list(self.h010(res)), ["BIP LITE BLACK"])
        self.assertEqual(self.reg(res, "0200")[0][2], "SMARTWATCH")
        self.assertTrue(any(c == "BIP LITE BLACK" and "confira" in m for t, c, m in res.pendencias if t == "item"))


class Sped(Base):
    Z = "|0000|019|0|01122025|31122025|LOJA X|11222333000181||CE|061234567|2304400||||A|1|\n"

    def test_0200_com_campos_vazios_no_fim(self):
        # Antes: o strip("|") comia os campos vazios e o 0200 era descartado (descrição perdida)
        txt = self.Z + "|0200|A|CANETA AZUL|||UN|00|\n|H005|31122025|20,00|01|\n|H010|A|UN|10,000|2,000000|20,00|0||||20,00|\n"
        res = self.gerar(arquivo("sped0200.txt", txt))
        self.assertEqual(self.reg(res, "0200")[0][2], "CANETA AZUL")
        self.assertFalse(self.pend(res, "sem descrição"))

    def test_campos_do_h010_mantidos(self):
        txt = (self.Z + "|0150|FORN1|FORNECEDOR|1058|11444777000161|||2304400||RUA|1||BAIRRO|\n"
               "|0200|A|CANETA|||UN|00|96081000|||||\n|H005|31122025|20,00|01|\n"
               "|H010|A|UN|10,000|2,000000|20,00|1|FORN1|EM CONSIGNACAO|1.1.3|15,00|\n")
        res = self.gerar(arquivo("spedterc.txt", txt))
        self.assertEqual(self.reg(res, "H010")[0][6:], ["1", "FORN1", "EM CONSIGNACAO", "1.1.3", "15,00"])
        self.assertEqual(self.reg(res, "0150")[0][1:3], ["FORN1", "FORNECEDOR"])
        self.assertFalse(self.pend(res, "COD_PART"))

    def test_participante_estrangeiro_sem_municipio(self):
        txt = ("|0150|EXT1|DISTRIBUIDOR EXTERIOR|5860||||||RUA X|1|||\n|0200|A|CANETA|||UN|00|96081000|||||\n"
               "|H005|31122025|20,00|01|\n|H010|A|UN|10,000|2,000000|20,00|1|EXT1|||20,00|\n")
        res = self.gerar(arquivo("sped_ext.txt", txt))
        self.assertEqual(self.reg(res, "0150")[0][1:4], ["EXT1", "DISTRIBUIDOR EXTERIOR", "5860"])
        self.assertEqual(self.reg(res, "H010")[0][7], "EXT1")

    def test_participante_sem_dados_tira_o_item(self):
        txt = ("|0200|A|CANETA|||UN|00|96081000|||||\n|H005|31122025|20,00|01|\n"
               "|H010|A|UN|10,000|2,000000|20,00|0||||20,00|\n|H010|B|UN|1,000|5,000000|5,00|2|SEM0150|||5,00|\n")
        res = self.gerar(arquivo("sped_sem0150.txt", txt))
        self.assertEqual(list(self.h010(res)), ["A"])
        self.assertTrue(self.pend(res, "participante SEM0150 sem nome e município"))

    def test_dois_inventarios_no_arquivo(self):
        # Antes: os H010 dos dois H005 eram somados num só inventário
        txt = ("|0200|A|CANETA|||UN|00|96081000|||||\n|H005|31122025|20,00|01|\n|H010|A|UN|10,000|2,000000|20,00|0||||20,00|\n"
               "|H005|31122025|30,00|05|\n|H010|A|UN|15,000|2,000000|30,00|0||||30,00|\n")
        res = self.gerar(arquivo("sped2h005.txt", txt))
        self.assertEqual(self.h010(res)["A"][3], "10,000")
        self.assertEqual(res.total, Decimal("20.00"))
        self.assertTrue(self.pend(res, "o arquivo tem 2 inventários"))

    def test_sped_com_bom(self):
        txt = "﻿|0200|A|CANETA|||UN|00|96081000|||||\n|H005|31122025|20,00|01|\n|H010|A|UN|10,000|2,000000|20,00|0||||20,00|\n"
        p = arquivo("sped_bom.txt", txt.encode("utf-8"))
        self.assertEqual(gs.ler_arquivo(p).tipo, "SPED / bloco H")


class Fortes(Base):
    def test_indicador_desconhecido_e_avisado(self):
        fs = ("CAB|001|1|FORTES|LOJA|11222333000181|31122025|\nPRO|1|ITEM|A1|96081000|UN|||001|001|\n"
              "IIV|1|31122025|1|1,000|UN|2,0000|09|2,00|\n")
        res = self.gerar(arquivo("prop.FS", fs))
        self.assertTrue(self.pend(res, "indicador de propriedade \"09\" desconhecido"))

    def test_item_de_terceiros_sem_participante_nao_invalida_o_arquivo(self):
        fs = ("CAB|001|1|FORTES|LOJA|11222333000181|31122025|\nPRO|1|ITEM|A1|96081000|UN|||001|001|\n"
              "PRO|2|ITEM2|A2|96081000|UN|||001|001|\n"
              "IIV|1|31122025|1|1,000|UN|2,0000|01|2,00|\nIIV|1|31122025|2|1,000|UN|3,0000|03|3,00|\n")
        res = self.gerar(arquivo("terc.FS", fs))
        self.assertEqual(list(self.h010(res)), ["A1"])
        self.assertEqual([i.codigo for i in res.fora], ["A2"])


class Texto(Base):
    def test_caracteres_do_word_viram_equivalentes(self):
        p = arquivo("word.csv", "Codigo;Descricao;Qtd;Valor Unit\nA1;CANETA – AZUL “PREMIUM” 5€ …;10;2,00\n")
        res = self.gerar(p)
        self.assertEqual(self.reg(res, "0200")[0][2], 'CANETA - AZUL "PREMIUM" 5EUR ...')
        self.assertFalse(self.pend(res, "ISO-8859-1"))

    def test_caractere_sem_equivalente_e_avisado(self):
        p = arquivo("chines.csv", "Codigo;Descricao;Qtd;Valor Unit\nA1;CANETA 笔 ŝ;10;2,00\n")
        res = self.gerar(p)
        self.assertEqual(self.reg(res, "0200")[0][2], "CANETA ? s")
        self.assertTrue(self.pend(res, "ISO-8859-1"))

    def test_controle_c1_e_acento_decomposto(self):
        self.assertEqual(gs.texto("A\x85B\x9fC"), "A B C")
        self.assertEqual(gs.texto("café"), "café")  # Mac grava o acento separado
        self.assertEqual(gs.texto(0), "0")  # antes: o código 0 sumia

    def test_codigo_alterado_e_avisado(self):
        p = arquivo("codigo.csv", "Codigo;Descricao;Qtd;Valor Unit\nA–1;X;1;1,00\nB|2;Y;1;1,00\n")
        res = self.gerar(p)
        self.assertEqual(sorted(self.h010(res)), ["A-1", "B 2"])
        self.assertEqual(len(self.pend(res, "caracteres que o SPED não aceita")), 2)

    def test_unidade_comprida_e_avisada(self):
        p = arquivo("unid.csv", "Codigo;Descricao;Unidade;Qtd;Valor Unit\nA1;X;CAIXA12;1;1,00\nA2;Y;Unidade;1;1,00\n")
        res = self.gerar(p)
        h = self.h010(res)
        self.assertEqual((h["A1"][2], h["A2"][2]), ("CAIXA1", "UN"))
        self.assertEqual(len(self.pend(res, "mais de 6 caracteres")), 1)

    def test_arquivo_gravado_em_iso_8859_1_e_crlf(self):
        p = arquivo("acentos.csv", "Codigo;Descricao;Qtd;Valor Unit\nA1;CALÇA AÇÃO – “X”;1;1,00\n")
        res = self.gerar(p)
        sped, _, _ = gs.gravar(res, TMP / "saida-iso")
        bruto = sped.read_bytes()
        self.assertIn('CALÇA AÇÃO - "X"'.encode("latin-1"), bruto)
        self.assertNotIn(b"?", bruto)
        self.assertTrue(bruto.endswith(b"|\r\n"))
        self.assertEqual(bruto.count(b"\n"), bruto.count(b"\r\n"))


class Cadastro(Base):
    def test_tipo_do_item_invalido(self):
        p = arquivo("tipo.csv", "Codigo;Descricao;Qtd;Valor Unit;Tipo\nA1;X;10;2,00;100\nA2;Y;1;1,00;7\nA3;Z;1;1,00;99\n")
        res = self.gerar(p, tipo_item="00")
        tipos = {c[1]: c[6] for c in self.reg(res, "0200")}
        self.assertEqual(tipos, {"A1": "00", "A2": "07", "A3": "99"})
        self.assertTrue(self.pend(res, "tipo do item \"100\""))

    def test_cest_invalido(self):
        p = arquivo("cest.csv", "Codigo;Descricao;Qtd;Valor Unit;CEST\nA1;X;10;2,00;123\nA2;Y;1;1,00;28.036.00\n")
        res = self.gerar(p)
        cest = {c[1]: c[12] for c in self.reg(res, "0200")}
        self.assertEqual(cest, {"A1": "", "A2": "2803600"})
        self.assertTrue(self.pend(res, "CEST \"123\""))

    def test_itens_do_complemento_fora_do_inventario_sao_avisados(self):
        inv = arquivo("inv.csv", "Codigo;Descricao;Qtd;Valor Unit\nA1;X;1;1,00\n")
        esc = arquivo("escriturado.csv", "Codigo;Descricao;Qtd\nA1;X;1\nZ9;OUTRO;5\n")
        res = self.gerar(inv, esc)
        self.assertTrue(self.pend(res, "1 item(ns) com quantidade neste arquivo não estão no inventário"))
        self.assertTrue(self.pend(res, "Z9"))

    def test_icms_somado_quando_codigos_iguais_se_juntam(self):
        p = arquivo("dup.csv", "Codigo;Descricao;Qtd;Valor Unit;ICMS Recuperavel\nA1;X;1;10,00;1,00\nA1;X;1;10,00;2,00\n")
        self.assertEqual(self.h010(self.gerar(p))["A1"][10], "17,00")


class Fiscal(Base):
    CSV = "Codigo;Descricao;Qtd;Valor Unit\nE1;X;1;1,00\n"

    def test_periodo_padrao_e_o_segundo_mes_depois_do_inventario(self):
        res = self.gerar(arquivo("p1.csv", self.CSV))
        self.assertEqual(self.reg(res, "0000")[0][1:5], ["020", "0", "01022026", "28022026"])
        self.assertFalse(self.pend(res, "Guia Prático"))
        self.assertEqual(gs.segundo_mes_depois(date(2025, 11, 30)), (1, 2026))
        self.assertEqual(gs.segundo_mes_depois(date(2025, 6, 30)), (8, 2025))

    def test_outro_periodo_e_avisado(self):
        res = self.gerar(arquivo("p2.csv", self.CSV), periodo="12/2025")
        self.assertEqual(self.reg(res, "0000")[0][1], "019")
        self.assertTrue(self.pend(res, "vai na EFD de 02/2026", "periodo"))

    def test_motivo_diferente_de_01_e_recusado(self):
        with self.assertRaisesRegex(gs.ErroConversao, "H020"):
            self.gerar(arquivo("p3.csv", self.CSV), mot_inv="02")

    def test_versao_do_leiaute(self):
        self.assertEqual(self.reg(self.gerar(arquivo("p4.csv", self.CSV), data_inv="31/12/2018"), "0000")[0][1], "013")
        res = self.gerar(arquivo("p5.csv", self.CSV), data_inv="31/12/2027")
        self.assertEqual(self.reg(res, "0000")[0][1], "020")
        self.assertTrue(self.pend(res, "não conhece o leiaute de 2028"))
        with self.assertRaisesRegex(gs.ErroConversao, "cod-ver"):
            self.gerar(arquivo("p6.csv", self.CSV), data_inv="31/12/2015")
        self.assertEqual(self.reg(self.gerar(arquivo("p7.csv", self.CSV), data_inv="31/12/2015", cod_ver="009"),
                                  "0000")[0][1], "009")

    def test_endereco_do_0005_faltando_e_avisado(self):
        res = self.gerar(arquivo("p8.csv", self.CSV), cnpj="11222333000181", uf="CE", cod_mun="2304400", ie="1", nome="X")
        for falta in ("CEP", "endereço", "bairro"):
            self.assertTrue(self.pend(res, f"falta {falta}", "empresa"), falta)
        res = self.gerar(arquivo("p9.csv", self.CSV), cnpj="11222333000181", uf="CE", cod_mun="2304400", ie="1", nome="X",
                         cep="60.000-000", end="RUA A", num="10", bairro="CENTRO")
        self.assertFalse(self.pend(res, "", "empresa"))
        self.assertEqual(self.reg(res, "0005")[0][2:7], ["60000000", "RUA A", "10", "", "CENTRO"])

    def test_municipio_de_outra_uf_e_avisado(self):
        res = self.gerar(arquivo("p10.csv", self.CSV), cnpj="11222333000181", uf="SP", cod_mun="2304400")
        self.assertTrue(self.pend(res, "não é um código IBGE de 7 dígitos da UF SP"))
        res = self.gerar(arquivo("p11.csv", self.CSV), uf="XX", cod_mun="123")
        self.assertTrue(self.pend(res, "UF \"XX\" não existe"))

    def test_escritorio_sem_crc_e_cpf_invalido(self):
        with mock.patch.object(gs, "ESCRITORIO", TMP / "esc.json"):
            (TMP / "esc.json").write_text('{"nome": "CONTADOR", "cpf": "111.444.777-35"}', encoding="utf-8")
            res = self.gerar(arquivo("p12.csv", self.CSV))
            self.assertEqual(self.reg(res, "0100")[0][1:4], ["CONTADOR", "11144477735", ""])
            self.assertTrue(self.pend(res, "falta o CRC"))
            self.assertFalse(self.pend(res, "CPF do contabilista"))
            (TMP / "esc.json").write_text('{"nome": "CONTADOR", "cpf": "111.444.777-00", "crc": "CE-1"}', encoding="utf-8")
            self.assertTrue(self.pend(self.gerar(arquivo("p13.csv", self.CSV)), "CPF do contabilista"))

    def test_cpf_valido(self):
        self.assertTrue(gs.cpf_valido("111.444.777-35"))
        self.assertFalse(gs.cpf_valido("111.444.777-36"))
        self.assertFalse(gs.cpf_valido("111.111.111-11"))


class Pendencias(Base):
    def test_csv_de_pendencias_nao_executa_formula(self):
        p = arquivo("formula.csv", 'Codigo;Descricao;Qtd;Valor Unit\n=HYPERLINK("http://x");X;-1;2,00\n+SOMA;Y;1;1,00\n'
                                   '12"3;Z;0;1,00\n')
        res = self.gerar(p)
        _, pend, _ = gs.gravar(res, TMP / "saida-formula")
        texto = pend.read_text(encoding="utf-8-sig")
        self.assertIn("'=HYPERLINK", texto)
        self.assertNotIn(";=HYPERLINK", texto)
        self.assertNotIn('"=HYPERLINK', texto)
        self.assertIn('"=""12""""3"""', texto)  # ="12""3" dentro do CSV

    def test_pendencia_fica_no_item_certo(self):
        # Código repetido com outra descrição é renomeado: o aviso do segundo item ia para o primeiro
        p = arquivo("renomeado.csv", "Codigo;Descricao;Qtd;Valor Unit;Valor Total\nA1;UM;1;1,00;1,00\nA1;DOIS;2;1,00;5,00\n")
        res = self.gerar(p)
        por = {i.codigo: i.pendencias for i in res.itens}
        self.assertEqual(por["A1"], [])
        self.assertTrue(any("total do arquivo 5,00" in m for m in por["A1-2"]))
        self.assertTrue(any(c == "A1-2" and "total do arquivo" in m for _, c, m in res.pendencias))


class Conferencias(Base):
    """O que o arquivo não traz e o app teve de supor vira aviso."""

    def test_quantidade_fracionada_em_unidade_inteira(self):
        p = arquivo("frac.csv", "Codigo;Descricao;Unidade;Qtd;Valor Unit\nA1;CANETA;UN;755,339;1,00\nA2;TECIDO;M;2,5;1,00\n"
                                "A3;LAPIS;PC;3;1,00\n")
        res = self.gerar(p)
        avisos = [(c, m) for t, c, m in res.pendencias if "fracionada" in m]
        self.assertEqual([c for c, _ in avisos], ["A1"])
        self.assertIn("755,339", avisos[0][1])

    def test_ncm_vazio_no_comercio_vira_um_aviso_so(self):
        p = arquivo("semncm.csv", "Codigo;Descricao;Qtd;Valor Unit;NCM\nA1;X;1;1,00;\nA2;Y;1;1,00;\nA3;Z;1;1,00;96081000\n")
        res = self.gerar(p)
        self.assertEqual(len(self.pend(res, "2 de 3 itens sem NCM", "arquivo")), 1)
        self.assertFalse(self.pend(res, "sem NCM:", "item"))

    def test_ncm_vazio_na_industria_avisa_cada_item(self):
        p = arquivo("semncm_ind.csv", "Codigo;Descricao;Qtd;Valor Unit;NCM\nA1;X;1;1,00;\nA3;Z;1;1,00;96081000\n")
        res = self.gerar(p, ind_ativ="0")
        self.assertEqual([c for t, c, m in res.pendencias if t == "item" and "sem NCM" in m], ["A1"])
        self.assertTrue(self.pend(res, "matéria-prima é 01", "arquivo"))

    def test_tipo_padrao_avisado(self):
        p = arquivo("semtipo.csv", "Codigo;Descricao;Qtd;Valor Unit;Tipo\nA1;X;1;1,00;\nA2;Y;1;1,00;04\n")
        # no comércio o tipo padrão 00 (revenda) é o certo: sem aviso
        self.assertFalse(self.pend(self.gerar(p), "sem o tipo no arquivo"))
        res = self.gerar(p, ind_ativ="0")
        self.assertEqual(len(self.pend(res, "1 de 2 itens sem o tipo no arquivo: usado o tipo padrão 00", "arquivo")), 1)
        # tipo de indústria escolhido numa empresa de comércio: o aviso diz o que fazer
        res = self.gerar(p, tipo_item="04")
        self.assertTrue(self.pend(res, "O tipo 04 (Produto acabado) é de indústria, mas a Atividade está como Outros", "arquivo"))

    def test_cnae_de_industria_marca_a_atividade(self):
        self.assertTrue(gs.cnae_industrial("1412601"))  # confecção de roupas
        self.assertTrue(gs.cnae_industrial("0500301"))
        self.assertFalse(gs.cnae_industrial("4783102"))  # comércio de relógios
        self.assertFalse(gs.cnae_industrial(""))
        consulta = {"nome": "MALHARIA X", "uf": "CE", "cod_mun": "2307650", "cep": "61900000", "end": "RUA A", "num": "1",
                    "compl": "", "bairro": "CENTRO", "fantasia": "", "fone": "", "email": "", "cnae": "1412601"}
        p = arquivo("cnae.csv", "Codigo;Descricao;Qtd;Valor Unit\nA1;X;1;1,00\n")
        f = gs.ler_arquivo(p)
        with mock.patch.object(gs, "consultar_cnpj", return_value=consulta), \
             mock.patch.object(gs, "CADASTRO_EMPRESAS", TMP / "empresas_cnae.json"):
            res = gs.converter(f, [], gs.opcoes(data_inv="31/12/2025", cnpj="71503824000122"))
            self.assertEqual(self.reg(res, "0000")[0][14], "0")
            self.assertEqual(len(self.reg(res, "0002")), 1)
            self.assertTrue(self.pend(res, "CNAE 1412601 é de indústria"))
            # escolhido no formulário, vale o que a pessoa marcou
            res = gs.converter(f, [], gs.opcoes(data_inv="31/12/2025", cnpj="71503824000122", ind_ativ="1"))
            self.assertEqual(self.reg(res, "0000")[0][14], "1")


class Rodada2(Base):
    """Melhorias dos testes manuais com os clientes fictícios."""

    CAB = "REGISTRO DE INVENTARIO\nFirma: LOJA\nEstoque existente em 31/12/2025\n"

    def test_total_diferente_de_quantidade_vezes_unitario(self):
        # Fórmula errada na planilha (=SOMA(qtd+unit)): o total era aceito sem aviso
        p = arquivo("formula.csv", "Codigo;Descricao;Qtd;Valor Unit;Valor Total\n10008;TECIDO;1232,89;12;1244,89\n"
                                   "1001642;COLETE;88;2;90\nA3;OK;3;2,47;7,40\nA4;OK2;10000;1,23;12345,00\n")
        res = self.gerar(p)
        h = self.h010(res)
        self.assertEqual((h["10008"][5], h["1001642"][5]), ("14794,68", "176,00"))
        self.assertEqual((h["A3"][5], h["A4"][5]), ("7,40", "12345,00"))  # unitário arredondado no arquivo: confere
        self.assertTrue(any(c == "10008" and "total do arquivo 1244,89" in m for _, c, m in res.pendencias))
        self.assertTrue(self.pend(res, "2 de 4 itens com total diferente de quantidade × unitário", "arquivo"))
        self.assertTrue(gs.total_confere(Decimal("3"), Decimal("2.47"), Decimal("7.40")))
        self.assertFalse(gs.total_confere(Decimal("88"), Decimal("2"), Decimal("90")))
        # unitário inteiro é o valor digitado, não um arredondamento: 1308 × 1 não é 1310
        self.assertFalse(gs.total_confere(Decimal("1308"), Decimal("1"), Decimal("1310")))
        self.assertTrue(gs.total_confere(Decimal("10000"), Decimal("1.23"), Decimal("12345")))

    def test_complemento_de_outro_estabelecimento(self):
        # Matriz junto com a filial: a matriz não pode completar quantidade nem dados da filial
        filial = arquivo("filial.txt", "REGISTRO DE INVENTARIO\nFirma: LOJA FILIAL\nInsc. Est.: 123  CNPJ: 11.222.333/0002-62\n"
                         "Estoque existente em 31/12/2025\n94036000 A1 CADEIRA          UN   2,000   10,0000   20,00\n")
        matriz = arquivo("matriz.txt", "REGISTRO DE INVENTARIO\nFirma: LOJA MATRIZ\nInsc. Est.: 999  CNPJ: 11.222.333/0001-81\n"
                         "Estoque existente em 31/12/2025\n94036000 A1 CADEIRA          UN   9,000   10,0000   90,00\n"
                         "94036000 B2 MESA             UN   1,000   50,0000   50,00\n")
        fontes = [gs.ler_arquivo(filial), gs.ler_arquivo(matriz)]
        res = gs.converter(fontes[0], [fontes[1]], gs.opcoes(sem_consulta=True))
        validar(self, res.registros)
        self.assertEqual(res.emp["cnpj"], "11222333000262")
        self.assertEqual(res.emp["ie"], "123")
        self.assertTrue(self.pend(res, "é de outro estabelecimento (CNPJ 11222333000181)", "arquivo"))
        self.assertFalse(self.pend(res, "difere da de"))
        self.assertFalse(self.pend(res, "não estão no inventário"))

    def test_complemento_de_outra_data_so_completa_o_cadastro(self):
        inv = arquivo("inv2025.csv", "Codigo;Qtd;Valor Unit\nA1;5;2,00\nA2;1;3,00\n")
        ano_passado = arquivo("inv2024.csv", "Inventário em 31/12/2024\nCodigo;Descricao;NCM;Qtd;Valor Unit\n"
                              "A1;CANETA;96081000;99;8,00\nZ9;OUTRO;96081000;7;1,00\n")
        res = self.gerar(inv, ano_passado)
        h = self.h010(res)
        self.assertEqual((h["A1"][3], h["A1"][4]), ("5,000", "2,000000"))  # quantidade e valor do inventário
        self.assertEqual(self.reg(res, "0200")[0][2:8:5], ["CANETA", "96081000"])  # descrição e NCM do cadastro
        self.assertTrue(self.pend(res, "é de outra data (31/12/2024)", "arquivo"))
        self.assertFalse(self.pend(res, "difere da de"))
        self.assertFalse(self.pend(res, "não estão no inventário"))

    def test_cadastro_do_fs_sem_inventario_completa(self):
        fs = ("CAB|001|1|FORTES|LOJA|11222333000181|31122024|\nGRP|1|ROUPA|4|04|\n"
              "PRO|1|BLUSA BASICA|B1|61061000|PC|||001|001|\nPRO|2|CALCA JEANS|C2|62046200|PC|||001|001|\n"
              "IIV|1|31122024|1|5,000|PC|10,0000|01|50,00|\n")
        inv = arquivo("bloco_h.txt", "|H005|31122025|30,00|01|\n|H010|C2|PC|1,000|30,000000|30,00|0||||30,00|\n")
        res = self.gerar(inv, arquivo("cad.FS", fs))
        self.assertEqual(self.reg(res, "0200")[0][2:8], ["CALCA JEANS", "", "", "PC", "04", "62046200"])

    def test_descricao_cortada_no_relatorio(self):
        linhas = "".join(f"9113{n:04d} X{n:03d} PULSEIRA COURO TRABALHADO GENUI  UN   1,000   1,0000   1,00\n" for n in range(4))
        f = gs.ler_arquivo(arquivo("cortado.txt", self.CAB + linhas))
        self.assertTrue(any("4 item(ns) com código e descrição no limite de 36 caracteres" in a for a in f.avisos))
        # junto com um cadastro, vale a descrição completa
        cad = arquivo("cad_cortado.csv", "Codigo;Descricao\nX000;PULSEIRA COURO TRABALHADO GENUINO MARROM\n")
        res = self.gerar(arquivo("cortado2.txt", self.CAB + linhas), cad)
        self.assertEqual(self.reg(res, "0200")[0][2], "PULSEIRA COURO TRABALHADO GENUINO MARROM")

    def test_codigo_do_traco_com_varias_descricoes(self):
        rel = self.CAB + ("85176277 RUN PRO BLUE - SMARTWATCH AURORA  UN   1,000   10,0000   10,00\n"
                          "85176277 RUN PRO BLUE - SMARTWATCH CRONOS  UN   1,000   10,0000   10,00\n")
        res = self.gerar(arquivo("traco2.txt", rel))
        self.assertTrue(self.pend(res, "\"RUN PRO BLUE\" aparece com 2 descrições diferentes", "arquivo"))

    def test_pendencia_nao_se_repete_quando_linhas_somam(self):
        rel = self.CAB + ("85176277 ZEN LITE - SMARTWATCH TEMPUS  UN   1,000   10,0000   10,00\n"
                          "85176277 ZEN LITE - SMARTWATCH TEMPUS  UN   2,000   10,0000   20,00\n")
        res = self.gerar(arquivo("soma.txt", rel))
        msgs = [m for t, c, m in res.pendencias if c == "ZEN LITE"]
        self.assertEqual(len(msgs), len(set(msgs)))

    def test_arquivo_sem_nenhuma_descricao_cada_item_fica_com_pendencia(self):
        txt = "|H005|31122025|30,00|01|\n|H010|A|UN|1,000|10,000000|10,00|0||||10,00|\n|H010|B|UN|2,000|10,000000|20,00|0||||20,00|\n"
        res = self.gerar(arquivo("sem0200.txt", txt))
        self.assertTrue(self.pend(res, "nenhum dos 2 itens tem descrição", "arquivo"))  # o que juntar
        self.assertEqual({c for t, c, m in res.pendencias if t == "item" and "sem descrição" in m}, {"A", "B"})
        self.assertTrue(all(it.pendencias and it.descricao == f"ITEM {it.codigo}" for it in res.itens))

    def test_tipo_de_industria_sugere_a_atividade(self):
        p = arquivo("tipos.csv", "Codigo;Descricao;Qtd;Valor Unit;Tipo\nA1;TECIDO;1;1,00;01\nA2;BLUSA;1;1,00;04\n")
        self.assertTrue(self.pend(self.gerar(p), "2 item(ns) com tipo de indústria", "arquivo"))
        self.assertFalse(self.pend(self.gerar(p, ind_ativ="0"), "tipo de indústria"))

    def test_inventario_fora_de_31_12(self):
        p = arquivo("outubro.csv", "Codigo;Descricao;Qtd;Valor Unit\nA1;X;1;1,00\n")
        self.assertTrue(self.pend(self.gerar(p, data_inv="31/10/2024"), "inventário em 31/10/2024 com o motivo 01", "arquivo"))
        self.assertFalse(self.pend(self.gerar(p, data_inv="31/12/2024"), "com o motivo 01"))

    def test_unidades_desconhecidas_e_sinonimas(self):
        p = arquivo("unids.csv", "Codigo;Descricao;Unidade;Qtd;Valor Unit\nA1;X;PCT;1;1,00\nA2;Y;PCTE;1;1,00\nA3;Z;MI;1;1,00\n")
        res = self.gerar(p)
        self.assertTrue(self.pend(res, "unidade(s) MI fora da tabela", "arquivo"))
        self.assertTrue(self.pend(res, "as siglas PCT e PCTE são a mesma unidade (PACOTE)", "arquivo"))

    def test_mensagens_da_consulta_do_cnpj(self):
        import urllib.error
        nao_existe = urllib.error.HTTPError("u", 404, "Not Found", {}, None)
        self.assertEqual(gs.erro_consulta("123", nao_existe), "CNPJ 123 não encontrado na Receita: confira o número")
        self.assertEqual(gs.erro_consulta("123", urllib.error.URLError("x")), "sem acesso à internet para consultar o CNPJ")

    def test_nome_da_receita_diferente_do_arquivo(self):
        self.assertTrue(gs.nomes_parecidos("RELOJOARIA PONTEIRO DE OURO - MATRIZ", "PONTEIRO DE OURO RELOGIOS LTDA"))
        self.assertFalse(gs.nomes_parecidos("RELOJOARIA PONTEIRO DE OURO - MATRIZ", "FULANO DE TAL"))
        consulta = {"nome": "FULANO DE TAL", "fantasia": "KLEBER JOIAS", "uf": "SP", "cod_mun": "3534302", "cep": "14620000",
                    "end": "RUA A", "num": "1", "compl": "", "bairro": "CENTRO", "fone": "", "email": "", "cnae": ""}
        p = arquivo("nome.txt", "REGISTRO DE INVENTARIO\nFirma: RELOJOARIA PONTEIRO DE OURO - MATRIZ\nCNPJ: 11.222.333/0001-81\n"
                    "Estoque existente em 31/12/2025\n91021110 10501 RELOGIO      UN   1,000   10,0000   10,00\n")
        with mock.patch.object(gs, "consultar_cnpj", return_value=consulta), \
             mock.patch.object(gs, "CADASTRO_EMPRESAS", TMP / "empresas_nome.json"):
            res = gs.converter(gs.ler_arquivo(p), [], gs.opcoes())
        self.assertTrue(self.pend(res, "não bate com o do CNPJ 11222333000181 na Receita", "empresa"))

    def test_mensagem_da_diferenca_do_total(self):
        p = arquivo("tot.csv", "Codigo;Descricao;Qtd;Valor Unit;Valor Total\nA1;X;1;1.000,00;1.000,00\n;;;Total:;1.000,01\n")
        res = self.gerar(p)
        self.assertTrue(self.pend(res, "total gerado 1.000,00 é 0,01 menor que o total do arquivo do cliente (1.000,01)", "total"))

    def test_csv_agrupado_por_item(self):
        p = arquivo("grupo_csv.csv", "Codigo;Descricao;Unidade;Qtd;Valor Unit;Valor Total\nA1;X;UN;1,5;1,00;9,00\nA2;Y;UN;1;1,00;1,00\n")
        res = self.gerar(p)
        _, pend, _ = gs.gravar(res, TMP / "saida-grupo")
        linhas = pend.read_text(encoding="utf-8-sig").splitlines()[1:]
        itens = [l.split(";")[1] for l in linhas if l.startswith("item;")]
        self.assertEqual(itens, sorted(itens, key=itens.index))  # as de A1 juntas
        self.assertEqual(len(itens), 2)
        primeira_item = next(n for n, l in enumerate(linhas) if l.startswith("item;"))
        self.assertTrue(all(not l.startswith("item;") for l in linhas[:primeira_item]))
        self.assertTrue(all(l.startswith("item;") for l in linhas[primeira_item:]))


class SpedCompleto(Base):
    def test_endereco_do_0005_e_sem_consulta_quando_o_arquivo_tem_tudo(self):
        txt = ("|0000|019|0|01122025|31122025|LOJA X LTDA|11222333000181||CE|061234567|2304400||||A|1|\n"
               "|0005|LOJA X|60000000|RUA A|10||CENTRO|8532000000||x@x.com|\n"
               "|0200|A|CANETA|||UN|00|96081000|||||\n|H005|31122025|10,00|01|\n|H010|A|UN|1,000|10,000000|10,00|0||||10,00|\n")
        f = gs.ler_arquivo(arquivo("completo.txt", txt))
        with mock.patch.object(gs, "consultar_cnpj", side_effect=AssertionError("não devia consultar")):
            res = gs.converter(f, [], gs.opcoes())
        validar(self, res.registros)
        self.assertEqual(self.reg(res, "0005")[0][1:8], ["LOJA X", "60000000", "RUA A", "10", "", "CENTRO", "8532000000"])
        self.assertFalse([p for p in res.pendencias if p[0] == "empresa"])


class CorrecoesDoPainel(Base):
    """Correções feitas no painel (args.ajustes) entram antes das conferências do app."""

    def planilha(self):
        return arquivo("estoque.csv", "Codigo;Descricao;Unidade;Quantidade;Valor unitario\n"
                                      "A1;CANETA AZUL;UN;10;2,50\n"
                                      "A2;LAPIS;UN;;1,00\n"
                                      ";BORRACHA;UN;3;0,90\n"
                                      ";BORRACHA;UN;5;0,90\n")

    def chaves(self, caminho):
        itens = gs.ler_arquivo(Path(caminho)).itens
        return [gs.chave_item(i) for i in itens]

    def test_chave_nao_muda_e_separa_linhas_sem_codigo(self):
        csv = self.planilha()
        a, b = self.chaves(csv), self.chaves(csv)
        self.assertEqual(a, b)
        self.assertEqual(len(set(a)), 4)  # as duas BORRACHA sem código são linhas diferentes

    def test_item_fora_entra_com_a_quantidade_corrigida(self):
        csv = self.planilha()
        lapis = self.chaves(csv)[1]
        res = self.gerar(csv)
        self.assertNotIn("A2", self.h010(res))
        self.assertEqual([i.codigo for i in res.fora], ["A2"])
        res = self.gerar(csv, ajustes={lapis: {"qtd": "1.000", "ncm": "96091000", "tipo": "7"}})
        self.assertEqual(self.h010(res)["A2"][3:6], ["1000,000", "1,000000", "1000,00"])
        lapis_final = next(i for i in res.itens if i.codigo == "A2")
        self.assertEqual((lapis_final.ncm, lapis_final.tipo, sorted(lapis_final.ajustado)), ("96091000", "07", ["ncm", "qtd", "tipo"]))
        self.assertEqual(res.fora, [])

    def test_codigo_corrigido(self):
        csv = self.planilha()
        chaves = self.chaves(csv)
        # Linha sem código ganha um código de verdade; o código antigo ainda acha o complemento
        cad = arquivo("cadastro.csv", "Codigo;Descricao;NCM\nA1;CANETA AZUL;96081000\n")
        res = self.gerar(csv, cad, ajustes={chaves[2]: {"codigo": "BOR-01"}, chaves[0]: {"codigo": "CAN-01"}})
        h = self.h010(res)
        self.assertIn("BOR-01", h)
        self.assertIn("CAN-01", h)
        self.assertNotIn("A1", h)
        caneta = next(i for i in res.itens if i.codigo == "CAN-01")
        self.assertEqual((caneta.ncm, caneta.ajustado), ("96081000", ["codigo"]))
        self.assertEqual([c[1] for c in self.reg(res, "0200") if c[1] == "CAN-01"], ["CAN-01"])
        # Código igual ao de outro item com outra descrição: o app não junta, renomeia e avisa
        res = self.gerar(csv, ajustes={chaves[2]: {"codigo": "A1"}})
        self.assertIn("A1-2", self.h010(res))
        self.assertTrue(self.pend(res, "repetido com outra descrição", "item"))

    def test_valor_corrigido_refaz_o_total(self):
        csv = self.planilha()
        caneta = self.chaves(csv)[0]
        res = self.gerar(csv, ajustes={caneta: {"vl_unit": "3,10", "descricao": "  ", "unidade": "cx"}})
        h = self.h010(res)["A1"]
        self.assertEqual((h[2], h[3:6]), ("CX", ["10,000", "3,100000", "31,00"]))
        self.assertEqual(next(i for i in res.itens if i.codigo == "A1").descricao, "CANETA AZUL")  # vazio não apaga


class ArquivosDeExemplo(Base):
    """Todo arquivo de exemplo da pasta arquivos/ gera um SPED que passa na conferência."""

    def test_todos_os_exemplos(self):
        casos = {
            "erp_titulo_e_total.xlsx": {}, "cod_produto_antigo.xls": {}, "loja_cp1252.csv": {},
            "relatorio_p7_paginas.txt": {}, "sped_completo.txt": {}, "segunda_aba_reais.xlsx": {},
            "ingles_utf8_bom.csv": {}, "INV2025.FS": {}, "relatorio_formato_real.txt": {},
            "fortes_layout_real.FS": {"data_inv": "31/12/2024"}, "modelo_preenchido.xlsx": {}, "ncm_numerico.xls": {},
            "relatorio_ncm_00.txt": {},
        }
        for nome, opc in casos.items():
            with self.subTest(nome):
                res = self.gerar(ARQ / nome, **opc)
                if res.principal.total_declarado is not None and nome != "relatorio_formato_real.txt":
                    self.assertEqual(res.total, res.principal.total_declarado)


if __name__ == "__main__":
    unittest.main()
