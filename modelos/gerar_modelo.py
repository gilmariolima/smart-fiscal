"""Gera o modelo de planilha de inventário para o cliente preencher (rode: python gerar_modelo.py)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

F = "Arial"
AZUL = "1453A3"; AMARELO = PatternFill("solid", fgColor="FFF2B3"); CINZA = PatternFill("solid", fgColor="EEF2F8")
fonte = lambda **k: Font(name=F, **k)
borda = Border(*(Side(style="thin", color="C9D3E3"),) * 4)
CAB = PatternFill("solid", fgColor=AZUL)

wb = Workbook()
ws = wb.active; ws.title = "Inventário"
LINHAS = 2000; INI = 9; FIM = INI + LINHAS - 1

# ---- cabeçalho do inventário (o app lê CNPJ e data daqui)
ws["A1"] = "INVENTÁRIO DE ESTOQUE"; ws["A1"].font = fonte(bold=True, size=14, color=AZUL)
campos = [("A2", "Empresa:", "B2"), ("A3", "CNPJ:", "B3"), ("A4", "Inventário em:", "B4")]
for rot, txt, val in campos:
    ws[rot] = txt; ws[rot].font = fonte(bold=True)
    ws[val].fill = AMARELO; ws[val].font = fonte(); ws[val].border = borda
ws.merge_cells("B2:D2"); ws.merge_cells("B3:C3")
ws["B3"].number_format = "@"
ws["B4"].number_format = "dd/mm/yyyy"
ws["B4"].comment = Comment("Data da contagem do estoque, normalmente 31/12. Digite como data: 31/12/2025.", "Estoque SPED")
ws["E4"] = "← data da contagem (ex.: 31/12/2025)"; ws["E4"].font = fonte(italic=True, color="6B7A90", size=9)
ws["E3"] = "← com ou sem pontos; aceita CNPJ com letras"; ws["E3"].font = fonte(italic=True, color="6B7A90", size=9)
ws["A5"] = "Total do inventário:"; ws["A5"].font = fonte(bold=True)
ws["B5"] = f"=SUM(I{INI}:I{FIM})"; ws["B5"].number_format = '#,##0.00'; ws["B5"].font = fonte(bold=True)
ws["A6"] = "Preencha as células em amarelo e uma linha por item a partir da linha 9. Colunas com * são obrigatórias. Veja a aba Instruções."
ws["A6"].font = fonte(italic=True, size=9, color="6B7A90")

# ---- tabela
cab = ["Código *", "Descrição", "Unidade", "NCM", "CEST", "Conta Contábil",
       "Quantidade *", "Valor Unitário *", "Valor Total", "ICMS Recuperável"]
larg = [21, 46, 10, 12, 10, 14, 13, 16, 15, 15]
for c, (t, w) in enumerate(zip(cab, larg), start=1):
    cel = ws.cell(INI - 1, c, t)
    cel.font = fonte(bold=True, color="FFFFFF"); cel.fill = CAB; cel.border = borda
    cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.column_dimensions[cel.column_letter].width = w
ws.row_dimensions[INI - 1].height = 30
dicas = {1: "Código do item no sistema do cliente. Texto: zeros à esquerda ficam.",
         3: "Sigla: UN, PC, KG, CX, MT, L... (por extenso também é aceito).",
         4: "8 dígitos, com ou sem pontos (ex.: 9403.60.00).",
         7: "Quantidade em estoque na data do inventário.",
         8: "Custo unitário do item (não o preço de venda).",
         9: "Calculado sozinho (Quantidade × Valor Unitário). Se só tiver o total, digite por cima.",
         10: "Opcional: ICMS a recuperar incluído no custo. Deixe 0 se não souber."}
for c, d in dicas.items():
    ws.cell(INI - 1, c).comment = Comment(d, "Estoque SPED")

for r in range(INI, FIM + 1):
    for c in range(1, 11):
        cel = ws.cell(r, c); cel.font = fonte(); cel.border = borda
    ws.cell(r, 1).number_format = "@"; ws.cell(r, 4).number_format = "@"; ws.cell(r, 5).number_format = "@"; ws.cell(r, 6).number_format = "@"
    ws.cell(r, 7).number_format = "#,##0.000"; ws.cell(r, 8).number_format = "#,##0.000000"
    ws.cell(r, 9).number_format = "#,##0.00"; ws.cell(r, 10).number_format = "#,##0.00"
    ws.cell(r, 9).value = f'=IF(OR(G{r}="",H{r}=""),"",ROUND(G{r}*H{r},2))'
    ws.cell(r, 9).fill = CINZA

num = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                     errorTitle="Valor inválido", error="Use um número maior ou igual a zero.")
ws.add_data_validation(num); num.add(f"G{INI}:H{FIM}"); num.add(f"J{INI}:J{FIM}")
dt = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True,
                    errorTitle="Data inválida", error="Digite a data como 31/12/2025.")
ws.add_data_validation(dt); dt.add("B4")
ws.freeze_panes = f"A{INI}"
ws.auto_filter.ref = f"A{INI-1}:J{FIM}"
ws.sheet_view.zoomScale = 100

# ---- instruções
ins = wb.create_sheet("Instruções")
ins.column_dimensions["A"].width = 22; ins.column_dimensions["B"].width = 16; ins.column_dimensions["C"].width = 70; ins.column_dimensions["D"].width = 24
ins["A1"] = "Como preencher o inventário"; ins["A1"].font = fonte(bold=True, size=14, color=AZUL)
ins["A2"] = "Com esta planilha preenchida, o Estoque SPED gera o bloco H do SPED Fiscal e o arquivo para importar no Domínio."
ins["A2"].font = fonte(size=10)
tab = [("Coluna", "Obrigatória?", "Como preencher", "Exemplo"),
       ("Código", "Sim", "Código do item no sistema da empresa.", "001003"),
       ("Quantidade", "Sim", "Quantidade em estoque na data do inventário.", "120"),
       ("Valor Unitário ou Valor Total", "Sim", "Custo do item. Com um dos dois, o outro é calculado. É o custo, não o preço de venda.", "7,20  ou  864,00"),
       ("Descrição", "Recomendada", "Nome do produto. Sem ela, sai \"ITEM <código>\".", "ÓLEO DE SOJA 900ML"),
       ("Unidade", "Recomendada", "Sigla da unidade. Sem ela, sai \"UN\".", "UN"),
       ("NCM", "Recomendada", "8 dígitos, com ou sem pontos.", "1507.90.11"),
       ("CEST", "Opcional", "7 dígitos, só para itens com substituição tributária.", "1700300"),
       ("Conta Contábil", "Opcional", "Conta do estoque na contabilidade.", "1.1.4.01"),
       ("ICMS Recuperável", "Opcional", "ICMS a recuperar incluído no custo. Deixe 0 se não souber.", "0")]
for i, linha in enumerate(tab, start=4):
    for j, v in enumerate(linha, start=1):
        cel = ins.cell(i, j, v); cel.border = borda; cel.alignment = Alignment(wrap_text=True, vertical="top")
        cel.font = fonte(bold=(i == 4), color="FFFFFF" if i == 4 else "000000")
        if i == 4: cel.fill = CAB
        elif linha[1] == "Sim": cel.fill = AMARELO
r = 4 + len(tab) + 1
ins.cell(r, 1, "Fora da tabela").font = fonte(bold=True, size=12, color=AZUL); r += 1
for t in ["• Data do inventário (célula B4 da aba Inventário). Se ficar em branco, o escritório digita no Estoque SPED.",
          "• Dados da empresa: nome e CNPJ (células B2 e B3). Inscrição estadual, UF e município o escritório informa uma vez no app.",
          "• Uma linha por item, sem linhas de subtotal no meio. Itens com quantidade zero ou valor negativo ficam fora do SPED.",
          "• Pode mandar a planilha no formato que o seu sistema exporta: as colunas acima com outros nomes (Cód. Produto, Referência, SKU, Qtde, Saldo, Custo Unitário, Vl. Unit.) também são reconhecidas."]:
    c = ins.cell(r, 1, t); c.font = fonte(size=10); ins.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    c.alignment = Alignment(wrap_text=True, vertical="top"); ins.row_dimensions[r].height = 28; r += 1
# ---- exemplo numa aba própria, com as mesmas colunas e larguras do inventário
exs = wb.create_sheet("Exemplo")
exs["A1"] = "Exemplo de preenchimento (não preencha esta aba: use a aba Inventário)"; exs["A1"].font = fonte(bold=True, size=12, color=AZUL)
for c, (t, w) in enumerate(zip(cab, larg), start=1):
    cel = exs.cell(3, c, t); cel.font = fonte(bold=True, color="FFFFFF"); cel.fill = CAB; cel.border = borda
    cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    exs.column_dimensions[cel.column_letter].width = w
exs.row_dimensions[3].height = 30
exemplos = [("001003", "ÓLEO DE SOJA 900ML", "UN", "1507.90.11", "", "1.1.4.01", 120, 7.2, None, 0, "com valor unitário"),
            ("A-77", "FEIJÃO CARIOCA 1KG", "PCT", "07133319", "", "1.1.4.01", 40, None, 271.56, 0, "só com o valor total (o unitário é calculado)")]
for r, ex in enumerate(exemplos, start=4):
    for c, v in enumerate(ex[:10], start=1):
        cel = exs.cell(r, c, v); cel.font = fonte(); cel.border = borda
    if ex[7] is not None:
        exs.cell(r, 9, f"=ROUND(G{r}*H{r},2)")
    exs.cell(r, 7).number_format = "#,##0.000"; exs.cell(r, 8).number_format = "#,##0.000000"; exs.cell(r, 9).number_format = "#,##0.00"
    exs.cell(r, 11, "← " + ex[10]).font = fonte(italic=True, size=9, color="6B7A90")
exs.column_dimensions["K"].width = 40
wb.active = 0
from pathlib import Path
wb.save(Path(__file__).with_name("Modelo_Inventario_Estoque_SPED.xlsx"))
print("ok")
